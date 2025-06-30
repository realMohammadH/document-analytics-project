import os
import re
import uuid
import time
from datetime import datetime
from flask import Flask, jsonify, request
from flask_cors import CORS
from werkzeug.utils import secure_filename
from io import BytesIO
from dotenv import load_dotenv
from supabase import create_client, Client
from postgrest.types import ReturnMethod

# Text extraction imports
try:
    import PyPDF2
except ImportError:
    PyPDF2 = None
try:
    import docx
except ImportError:
    docx = None

# Configuration
app = Flask(__name__)
CORS(app)
# CORS(app, resources={r"/api/*": {"origins": r"https://.*\.netlify\.app"}})

app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024  # 10MB file size limit

# Supabase Configuration
load_dotenv()

SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_KEY")

if not SUPABASE_URL or not SUPABASE_KEY:
    raise ValueError("SUPABASE_URL and SUPABASE_KEY must be set in environment variables")

supabase = create_client(SUPABASE_URL, SUPABASE_KEY)

total_search_count = 0


# Helper functions
def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in {
        "txt", "pdf", "doc", "docx"
    }

def highlight_search_terms(text, search_terms, highlight_tag="<mark>"):
    """Highlight search terms in text while preserving case sensitivity"""
    if not text or not search_terms:
        return text
    
    highlighted_text = text
    sorted_terms = sorted(search_terms, key=len, reverse=True)
    
    for term in sorted_terms:
        pattern = re.compile(r'\b' + re.escape(term) + r'\b', re.IGNORECASE)
        highlighted_text = pattern.sub(f"{highlight_tag}\\g<0></mark>", highlighted_text)
    
    return highlighted_text

def extract_text_snippets(text, search_terms, snippet_length=100):
    """Extract text snippets around search terms"""
    if not text or not search_terms:
        return []
    
    snippets = []
    text_lower = text.lower()
    
    for term in search_terms:
        term_lower = term.lower()
        start_pos = 0
        
        while True:
            pos = text_lower.find(term_lower, start_pos)
            if pos == -1:
                break
            
            snippet_start = max(0, pos - snippet_length // 2)
            snippet_end = min(len(text), pos + len(term) + snippet_length // 2)
            
            snippet = text[snippet_start:snippet_end]
            
            if snippet_start > 0:
                snippet = "..." + snippet
            if snippet_end < len(text):
                snippet = snippet + "..."
            
            highlighted_snippet = highlight_search_terms(snippet, [term])
            
            snippets.append({
                "term": term,
                "snippet": highlighted_snippet,
                "position": pos
            })
            
            start_pos = pos + 1
    
    # Remove duplicates and limit results
    unique_snippets = []
    seen_positions = set()
    
    for snippet in sorted(snippets, key=lambda x: x["position"]):
        is_duplicate = any(abs(snippet["position"] - seen_pos) < snippet_length 
                          for seen_pos in seen_positions)
        
        if not is_duplicate:
            unique_snippets.append(snippet)
            seen_positions.add(snippet["position"])
    
    return unique_snippets[:5]

def extract_text_from_file(file_path, file_type):
    """Extract text from PDF, DOCX, or TXT file"""
    text = ""
    
    if file_type == "pdf" and PyPDF2:
        try:
            with open(file_path, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text += page.extract_text() or ""
        except Exception:
            text = ""
    elif file_type == "docx" and docx:
        try:
            doc = docx.Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs])
        except Exception:
            text = ""
    elif file_type == "txt":
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        except Exception:
            text = ""
    
    return text.strip()

def extract_text_from_file_stream(file_stream, file_type):
    """Extract text from a file stream without saving to disk"""
    text = ""
    file_stream.seek(0)
    
    if file_type == "pdf" and PyPDF2:
        try:
            reader = PyPDF2.PdfReader(file_stream)
            for page in reader.pages:
                text += page.extract_text() or ""
        except Exception:
            text = ""
    elif file_type == "docx" and docx:
        try:
            doc = docx.Document(file_stream)
            text = "\n".join([para.text for para in doc.paragraphs])
        except Exception:
            text = ""
    elif file_type == "txt":
        try:
            text = file_stream.read().decode('utf-8', errors='ignore')
        except Exception:
            text = ""
    
    file_stream.seek(0)
    return text.strip()

def extract_title_from_content(content):
    """Extract the first meaningful line as title"""
    if not content:
        return "Untitled Document"
    
    for line in content.splitlines():
        line = line.strip()
        if line and len(line) > 5:
            return line[:120]
    
    return "Untitled Document"

# Classification system
CLASSIFICATION_TREE = {
    "Academic": [
        "research", "study", "analysis", "methodology", "hypothesis", "conclusion", 
        "literature", "review", "citation", "reference", "academic", "university", 
        "college", "education", "learning", "teaching", "student", "professor", 
        "lecture", "assignment", "thesis", "dissertation", "paper", "journal", 
        "publication", "scholarly", "scientific", "experiment", "data", "statistics",
        "survey", "questionnaire", "interview", "observation", "theory", "concept",
        "framework", "model", "approach", "method", "technique", "procedure"
    ],
    "IT": [
        "software", "hardware", "programming", "code", "algorithm", "database", 
        "system", "network", "server", "client", "application", "web", "mobile", 
        "cloud", "security", "cybersecurity", "encryption", "authentication", 
        "authorization", "firewall", "virus", "malware", "backup", "recovery", 
        "maintenance", "update", "patch", "version", "deployment", "testing", 
        "debugging", "optimization", "performance", "scalability", "architecture", 
        "framework", "api", "interface", "user", "experience", "design", "development",
        "agile", "scrum", "devops", "ci", "cd", "git", "repository", "branch", "merge"
    ],
    "Business": [
        "business", "company", "organization", "corporation", "enterprise", "startup", 
        "management", "leadership", "strategy", "planning", "marketing", "sales", 
        "finance", "accounting", "budget", "revenue", "profit", "loss", "investment", 
        "stock", "market", "customer", "client", "product", "service", "brand", 
        "advertising", "promotion", "campaign", "competitive", "analysis", 
        "market", "research", "survey", "feedback", "satisfaction", "loyalty", 
        "retention", "acquisition", "growth", "expansion", "partnership", "alliance", 
        "merger", "acquisition", "contract", "agreement", "negotiation", "deal", 
        "transaction", "payment", "invoice", "billing", "pricing", "cost", "expense",
        "roi", "kpi", "metric", "performance", "target", "goal", "objective", "mission",
        "vision", "values", "culture", "team", "employee", "staff", "workforce", "hr",
        "recruitment", "training", "development", "career", "promotion", "compensation"
    ]
}

def classify_document(content):
    """Classify document based on predefined categories"""
    if not content:
        return "Other"
    
    content_lower = content.lower()
    scores = {}
    
    for category, keywords in CLASSIFICATION_TREE.items():
        score = 0
        for keyword in keywords:
            matches = len(re.findall(r'\b' + re.escape(keyword) + r'\b', content_lower))
            score += matches
        scores[category] = score
    
    if not scores or all(s == 0 for s in scores.values()):
        return "Other"

    best_category = max(scores, key=lambda k: scores[k])
    best_score = scores[best_category]
    
    total_matches = sum(scores.values())
    confidence = best_score / total_matches if total_matches > 0 else 0
    
    return best_category if confidence >= 0.3 else "Other"

def extract_keywords(text, num_keywords=5):
    """Extract keywords using frequency analysis"""
    words = re.findall(r"\b[a-zA-Z]{3,}\b", text.lower())
    
    # Common stop words
    stop_words = {
        'the', 'be', 'to', 'of', 'and', 'a', 'in', 'that', 'have', 'i', 'it', 
        'for', 'not', 'on', 'with', 'he', 'as', 'you', 'do', 'at', 'this', 
        'but', 'his', 'by', 'from', 'document', 'file', 'will', 'are', 'can'
    }
    
    filtered_words = [word for word in words if word not in stop_words and len(word) > 3]
    
    word_freq = {}
    for word in filtered_words:
        word_freq[word] = word_freq.get(word, 0) + 1
    
    keywords = sorted(word_freq.items(), key=lambda x: x[1], reverse=True)
    return [word for word, freq in keywords[:num_keywords]]

# API Endpoints
@app.route("/api/health")
def health_check():
    return jsonify({
        "status": "healthy",
        "message": "Document Analytics System is running",
        "timestamp": datetime.utcnow().isoformat() + "Z",
        "version": "1.0.0"
    })

@app.route("/api/documents")
def get_documents():
    page = int(request.args.get("page", 1))
    per_page = min(int(request.args.get("per_page", 10)), 50)

    try:
        response = supabase.table('documents').select('*').order('upload_date', desc=True).execute()
        documents = response.data
        
        start = (page - 1) * per_page
        end = start + per_page
        
        return jsonify({
            "documents": documents[start:end],
            "pagination": {
                "page": page,
                "per_page": per_page,
                "total": len(documents),
                "pages": (len(documents) + per_page - 1) // per_page,
            },
        })
    except Exception as e:
        return jsonify({"error": "Could not fetch documents"}), 500

@app.route("/api/search")
def search_documents():
    start_time = time.time()
    global total_search_count
    
    query = request.args.get("q", "")
    if not query:
        return jsonify({"error": "Search query is required"}), 400
    
    total_search_count += 1
    query_terms = query.lower().split()
    
    # Store results for each keyword separately
    keyword_results = {}
    all_doc_ids = set()
    
    try:
        # Search for each keyword individually
        for term in query_terms:
            # Perform individual search for this keyword
            response = supabase.table('documents').select('*').text_search('content', term).execute()
            docs = response.data
            
            term_results = []
            for doc in docs:
                doc_id = doc.get("id")
                all_doc_ids.add(doc_id)
                
                # Get text snippets for this specific keyword
                text_snippets = extract_text_snippets(doc.get("content", ""), [term])
                
                term_results.append({
                    "id": doc_id,
                    "title": doc.get("title", ""),
                    "extracted_title": doc.get("extracted_title", ""),
                    "classification": doc.get("classification"),
                    "content_preview": doc.get("content_preview", ""),
                    "text_snippets": text_snippets,
                    "file_type": doc.get("file_type", ""),
                    "file_size": doc.get("file_size", 0),
                    "upload_date": doc.get("upload_date", "")
                })
            
            keyword_results[term] = {
                "keyword": term,
                "count": len(term_results),
                "documents": term_results
            }
        
        elapsed = time.time() - start_time
        
        return jsonify({
            "query": query,
            "query_terms": query_terms,
            "keyword_results": keyword_results,
            "total_unique_documents": len(all_doc_ids),
            "search_time_seconds": round(elapsed, 4)
        })
    except Exception as e:
        return jsonify({"error": f"Search failed: {str(e)}"}), 500

@app.route("/api/statistics")
def get_statistics():
    try:
        response = supabase.table('documents').select('classification, file_size, word_count').execute()
        docs = response.data

        total_docs = len(docs)
        if total_docs == 0:
            return jsonify({
                "documents": {"total": 0, "total_size_mb": 0, "total_words": 0},
                "classifications": [],
                "search": {"total_searches": total_search_count}
            })

        total_size_mb = sum(doc.get("file_size", 0) for doc in docs) / (1024 * 1024)
        total_words = sum(doc.get("word_count", 0) for doc in docs)

        classifications = {}
        for doc in docs:
            cat = doc["classification"]
            classifications[cat] = classifications.get(cat, 0) + 1

        classification_dist = [
            {
                "name": cat,
                "count": count,
                "percentage": round((count / total_docs) * 100, 2),
            }
            for cat, count in classifications.items()
            if cat
        ]

        return jsonify({
            "documents": {
                "total": total_docs,
                "total_size_mb": round(total_size_mb, 2),
                "total_words": total_words,
            },
            "classifications": classification_dist,
            "search": {"total_searches": total_search_count},
        })
    except Exception as e:
        return jsonify({"error": "Could not fetch statistics"}), 500

@app.route("/api/upload", methods=["POST"])
def upload_file():
    if "file" not in request.files:
        return jsonify({"error": "No file part"}), 400
    
    file = request.files["file"]
    if not file.filename or not allowed_file(file.filename):
        return jsonify({"error": "File type not allowed"}), 400

    filename = secure_filename(file.filename)
    if not filename:
        return jsonify({"error": "Invalid filename"}), 400

    file_type = filename.rsplit(".", 1)[1].lower()
    file_bytes = file.read()
    file_stream = BytesIO(file_bytes)

    content = extract_text_from_file_stream(file_stream, file_type)
    if not content:
        content = f"[No extractable text found in {filename}]"
    
    # Upload to Supabase Storage
    try:
        unique_filename = f"uploads/{uuid.uuid4()}_{filename}"
        supabase.storage.from_("documents").upload(
            file=file_bytes,
            path=unique_filename,
            file_options={"content-type": file.content_type or 'application/octet-stream'}
        )
        storage_url = supabase.storage.from_("documents").get_public_url(unique_filename)
    except Exception as e:
        print(f"Storage upload error: {str(e)}")
        print(f"Error type: {type(e)}")
        return jsonify({"error": f"Failed to upload file to storage: {str(e)}"}), 500

    # Process document
    extracted_title = extract_title_from_content(content)
    classification = classify_document(content)
    keywords = extract_keywords(content)
    
    new_doc_data = {
        "title": filename,
        "extracted_title": extracted_title,
        "file_type": file_type,
        "file_size": len(file_bytes),
        "classification": classification,
        "keywords": keywords,
        "upload_date": datetime.utcnow().isoformat() + "Z",
        "content_preview": content[:200],
        "word_count": len(content.split()),
        "reading_time": len(content.split()) // 200,
        "content": content,
        "storage_url": storage_url
    }
    
    try:
        response = supabase.table('documents').insert(new_doc_data, returning=ReturnMethod.representation).execute()
        
        if not response.data:
            return jsonify({"error": "Failed to save document"}), 500

        return jsonify({
            "success": True,
            "message": "File uploaded successfully",
            "document": response.data[0],
        })
    except Exception as e:
        return jsonify({"error": "Failed to save document metadata"}), 500

@app.route("/api/sorted_documents")
def get_sorted_documents():
    start_time = time.time()
    try:
        response = supabase.table('documents').select('*').order('extracted_title', desc=False).execute()
        elapsed = time.time() - start_time
        
        return jsonify({
            "documents": response.data,
            "sort_time_seconds": round(elapsed, 4),
            "total": len(response.data)
        })
    except Exception as e:
        return jsonify({"error": "Could not fetch sorted documents"}), 500

if __name__ == "__main__":
    app.run(debug=True, port=5000)
